#!/usr/bin/env python3
"""
DOCX to Markdown Converter
将 OneNET 城市物联网平台接口文档从 DOCX 转换为 Markdown
"""

import sys
import re
import subprocess

# 安装依赖
def install_package(package):
    subprocess.check_call([sys.executable, "-m", "pip", "install", package, "-q"])

try:
    from docx import Document
except ImportError:
    print("正在安装 python-docx...")
    install_package("python-docx")
    from docx import Document

try:
    from docx.table import Table
    from docx.text.paragraph import Paragraph
except ImportError:
    pass


def get_heading_level(paragraph):
    """获取段落的标题级别"""
    style_name = paragraph.style.name if paragraph.style else ""

    if "Heading 1" in style_name or "标题 1" in style_name:
        return 1
    elif "Heading 2" in style_name or "标题 2" in style_name:
        return 2
    elif "Heading 3" in style_name or "标题 3" in style_name:
        return 3
    elif "Heading 4" in style_name or "标题 4" in style_name:
        return 4
    elif "Heading 5" in style_name or "标题 5" in style_name:
        return 5
    elif "Heading" in style_name or "标题" in style_name:
        return 2

    # 基于文本模式判断
    text = paragraph.text.strip()
    if re.match(r'^\d+\.\s+\S', text):
        return 1
    elif re.match(r'^\d+\.\d+\.?\s+\S', text):
        return 2
    elif re.match(r'^\d+\.\d+\.\d+\.?\s+\S', text):
        return 3

    return 0


def table_to_markdown(table):
    """将 Word 表格转换为 Markdown 表格"""
    markdown_rows = []

    for i, row in enumerate(table.rows):
        cells = [cell.text.strip().replace('\n', ' ').replace('|', '\\|') for cell in row.cells]
        markdown_rows.append("| " + " | ".join(cells) + " |")

        # 在第一行后添加分隔符
        if i == 0:
            separator = "| " + " | ".join(["---"] * len(cells)) + " |"
            markdown_rows.append(separator)

    return "\n".join(markdown_rows)


def paragraph_to_markdown(paragraph):
    """将段落转换为 Markdown"""
    text = paragraph.text.strip()

    if not text:
        return ""

    heading_level = get_heading_level(paragraph)

    if heading_level > 0:
        return f"\n{'#' * heading_level} {text}\n"

    # 处理列表
    style_name = paragraph.style.name if paragraph.style else ""
    if "List" in style_name or "列表" in style_name:
        return f"- {text}"

    # 检查是否是代码块（简单启发式）
    if text.startswith('{') or text.startswith('[') or 'function' in text.lower():
        return f"```\n{text}\n```"

    return text


def convert_docx_to_markdown(docx_path, output_path):
    """将 DOCX 转换为 Markdown"""

    print(f"正在读取 DOCX 文件: {docx_path}")

    doc = Document(docx_path)
    markdown_lines = []

    markdown_lines.append("# OneNET 城市物联网平台 V3.0 公开接口文档\n")
    markdown_lines.append("> 本文档由 DOCX 自动转换生成\n")
    markdown_lines.append("---\n")

    # 遍历文档中的所有元素
    for element in doc.element.body:
        # 处理段落
        if element.tag.endswith('p'):
            for para in doc.paragraphs:
                if para._element == element:
                    md_text = paragraph_to_markdown(para)
                    if md_text:
                        markdown_lines.append(md_text)
                    break

        # 处理表格
        elif element.tag.endswith('tbl'):
            for table in doc.tables:
                if table._element == element:
                    markdown_lines.append("\n")
                    markdown_lines.append(table_to_markdown(table))
                    markdown_lines.append("\n")
                    break

    # 写入 Markdown 文件
    print(f"正在写入 Markdown 文件: {output_path}")

    with open(output_path, 'w', encoding='utf-8') as f:
        f.write("\n".join(markdown_lines))

    print(f"✓ 转换完成！")
    print(f"输出文件: {output_path}")


if __name__ == "__main__":
    docx_file = "./对接/副本11-OneNET城市物联网平台V3.0公开接口文档.docx"
    output_file = "./requirements/OneNET城市物联网平台V3.0公开接口文档-docx.md"

    convert_docx_to_markdown(docx_file, output_file)
